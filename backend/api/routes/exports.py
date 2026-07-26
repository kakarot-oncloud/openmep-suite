"""
CSV / Excel export endpoints.

Turns any tabular result (a BOQ, a cable schedule, a calc summary) into a
downloadable ``.csv`` or ``.xlsx`` file. Generic table endpoints accept
``{columns, rows}``; the cable-schedule convenience endpoint sizes a batch and
returns the spreadsheet in one call.
"""

import csv
import io
import re
from typing import Any

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from backend.models.project import CableScheduleBatchRequest, TableExportRequest
from backend.services.cable_schedule import size_schedule

router = APIRouter(prefix="/exports", tags=["Exports"])

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _safe_filename(name: str, ext: str) -> str:
    # Disallow dots in the stem so path-traversal fragments like ".." cannot survive.
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_") or "openmep_export"
    return f"{stem[:80]}.{ext}"


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return " | ".join(str(v) for v in value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def build_xlsx(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\\/*?:\[\]]", " ", title)[:31]) or "Export"

    ws.append(columns)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([_cell(v) for v in row])

    # Approximate column auto-width.
    for idx, col in enumerate(columns, start=1):
        longest = len(str(col))
        for row in rows:
            if idx - 1 < len(row):
                longest = max(longest, len(str(_cell(row[idx - 1]))))
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(longest + 2, 10), 60)

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(columns: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_cell(v) for v in row])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def _xlsx_response(title: str, columns: list[str], rows: list[list[Any]]) -> StreamingResponse:
    data = build_xlsx(title, columns, rows)
    fname = _safe_filename(title, "xlsx")
    return StreamingResponse(
        io.BytesIO(data),
        media_type=XLSX_MEDIA,
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


def _csv_response(title: str, columns: list[str], rows: list[list[Any]]) -> StreamingResponse:
    data = build_csv(columns, rows)
    fname = _safe_filename(title, "csv")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.post("/table.xlsx", summary="Export a table as an Excel workbook")
async def export_table_xlsx(req: TableExportRequest) -> Any:
    return _xlsx_response(req.title, req.columns, req.rows)


@router.post("/table.csv", summary="Export a table as CSV")
async def export_table_csv(req: TableExportRequest) -> Any:
    return _csv_response(req.title, req.columns, req.rows)


@router.post("/cable-schedule.xlsx", summary="Size a cable schedule and export it as Excel")
async def export_cable_schedule_xlsx(req: CableScheduleBatchRequest) -> Any:
    result = size_schedule(req.design_basis, req.circuits)
    table = result["table"]
    return _xlsx_response("Cable_Schedule", table["columns"], table["rows"])
