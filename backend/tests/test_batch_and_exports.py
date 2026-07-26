"""
Tests for batch cable-schedule sizing and CSV / Excel export endpoints.
"""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.main import app

client = TestClient(app)

BATCH_PAYLOAD = {
    "design_basis": {
        "region": "gcc",
        "sub_region": "dewa",
        "ambient_temp_c": 45,
        "power_factor": 0.85,
        "default_cable_type": "XLPE_CU",
        "default_installation_method": "C",
    },
    "circuits": [
        {"circuit_ref": "C1", "description": "AHU-1", "load_kw": 45, "cable_length_m": 80, "phases": 3},
        {"circuit_ref": "C2", "description": "Lighting DB", "load_kw": 6, "cable_length_m": 40,
         "phases": 1, "circuit_type": "lighting"},
        {"circuit_ref": "C3", "description": "Chiller", "load_kw": 160, "cable_length_m": 120, "phases": 3},
    ],
}


class TestBatchCableSchedule:
    def test_batch_sizes_all_circuits(self):
        r = client.post("/api/electrical/cable-schedule", json=BATCH_PAYLOAD)
        assert r.status_code == 200
        d = r.json()
        assert d["status"] == "success"
        assert d["circuit_count"] == 3
        assert len(d["circuits"]) == 3
        for c in d["circuits"]:
            assert c["selected_size_mm2"] > 0
            assert c["design_current_ib_a"] > 0
        # Table is export-ready
        assert d["table"]["columns"][0] == "Ref"
        assert len(d["table"]["rows"]) == 3

    def test_circuit_inherits_design_basis(self):
        """A circuit with no PF of its own uses the design-basis power factor."""
        r = client.post("/api/electrical/cable-schedule", json=BATCH_PAYLOAD).json()
        assert r["circuits"][0]["power_factor"] == 0.85

    def test_per_circuit_override_wins(self):
        payload = {
            "design_basis": {"region": "gcc", "power_factor": 0.85},
            "circuits": [{"load_kw": 20, "cable_length_m": 50, "power_factor": 0.95}],
        }
        r = client.post("/api/electrical/cable-schedule", json=payload).json()
        assert r["circuits"][0]["power_factor"] == 0.95

    def test_all_regions_batch(self):
        for region in ["gcc", "europe", "india", "australia"]:
            payload = {
                "design_basis": {"region": region},
                "circuits": [{"load_kw": 30, "cable_length_m": 60, "phases": 3}],
            }
            r = client.post("/api/electrical/cable-schedule", json=payload)
            assert r.status_code == 200, f"{region} failed: {r.text}"
            assert r.json()["circuits"][0]["selected_size_mm2"] > 0

    def test_empty_circuits_rejected(self):
        r = client.post("/api/electrical/cable-schedule",
                        json={"design_basis": {"region": "gcc"}, "circuits": []})
        assert r.status_code == 422


class TestExports:
    def test_table_csv(self):
        r = client.post("/api/exports/table.csv", json={
            "title": "My Schedule",
            "columns": ["A", "B"],
            "rows": [[1, "x"], [2, "y"]],
        })
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("text/csv")
        assert "attachment" in r.headers["content-disposition"]
        text = r.content.decode("utf-8-sig")
        assert "A,B" in text and "1,x" in text

    def test_table_xlsx(self):
        r = client.post("/api/exports/table.xlsx", json={
            "title": "My Schedule",
            "columns": ["A", "B"],
            "rows": [[1, "x"], [2, "y"]],
        })
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        assert [c.value for c in ws[1]] == ["A", "B"]
        assert ws.cell(row=2, column=1).value == 1

    def test_cable_schedule_xlsx(self):
        r = client.post("/api/exports/cable-schedule.xlsx", json=BATCH_PAYLOAD)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # header + 3 circuits
        assert ws.max_row == 4
        assert ws.cell(row=1, column=1).value == "Ref"

    def test_export_filename_sanitised(self):
        r = client.post("/api/exports/table.csv", json={
            "title": "../../etc/passwd\"; rm -rf",
            "columns": ["A"], "rows": [[1]],
        })
        assert r.status_code == 200
        cd = r.headers["content-disposition"]
        assert ".." not in cd and "/" not in cd.split("filename=")[1]
